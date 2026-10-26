import os
import pytest
import openpyxl
import pandas as pd
from app.engine.forecast import compute_forecast_from_df
from scripts.build_legacy_workbook import build_legacy_workbook
from scripts.generate_synthetic_data import create_snapshots, get_as_of_date, snapshot_filename

def todays_snapshot_path(base_dir: str) -> str:
    """Path of today's snapshot, generating the snapshot pair if it does not exist yet."""
    snap_dir = os.path.join(base_dir, "data", "snapshots")
    path = os.path.join(snap_dir, snapshot_filename(get_as_of_date()))
    if not os.path.exists(path):
        create_snapshots(count=2000, output_dir=snap_dir)
    return path

def load_legacy_workbook_pivot_values(wb_path: str):
    """
    Reads and evaluates the formulas and static cells directly from the Pivot sheet
    of Daily_Forecast_LEGACY.xlsx based on the underlying RawExtract sheet.
    """
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    ws_pivot = wb["Pivot"]
    ws_raw = wb["RawExtract"]

    # 1. Read RawExtract into DataFrame
    raw_data = []
    headers = [cell.value for cell in ws_raw[1]]
    for row in ws_raw.iter_rows(min_row=2, values_only=True):
        raw_data.append(dict(zip(headers, row)))
    df_raw = pd.DataFrame(raw_data)
    df_raw["saas_acv_usd_at_plan"] = pd.to_numeric(df_raw["saas_acv_usd_at_plan"], errors="coerce").fillna(0.0)

    # 2. Read filter cell values from Pivot sheet
    filter_country = ws_pivot["C6"].value
    filter_client_sub = ws_pivot["C7"].value
    filter_signing_cat = ws_pivot["C8"].value
    filter_ut15 = ws_pivot["C9"].value

    # Base filter mask
    mask = (
        (df_raw["country"] == filter_country) &
        (df_raw["client_sub_type"] == filter_client_sub) &
        (df_raw["signing_category"] == filter_signing_cat) &
        (df_raw["ut_level_15_dscr"] == filter_ut15)
    )

    # Check for hardcoded exclusion in formula text (e.g. <>Acme Test Enterprise 999)
    formula_sample = str(ws_pivot["C12"].value or "")
    if "<>Acme Test Enterprise 999" in formula_sample:
        mask = mask & (df_raw["customer_name"] != "Acme Test Enterprise 999")

    df_filtered = df_raw[mask].copy()

    # 3. Read per-product rows (rows 12-17)
    ut30_values = {}
    for r in range(12, 18):
        product_name = ws_pivot.cell(row=r, column=2).value
        prod_df = df_filtered[df_filtered["ut_level_30_dscr"] == product_name]
        deal_count = int(len(prod_df))
        total_acv = round(float(prod_df["saas_acv_usd_at_plan"].sum()), 2)
        avg_acv = round(total_acv / deal_count, 2) if deal_count > 0 else 0.0
        
        ut30_values[product_name] = {
            "deal_count": deal_count,
            "total_acv": total_acv,
            "avg_acv": avg_acv
        }

    # 4. Read Manual Adjustment (row 18)
    manual_adj = {
        "description": ws_pivot.cell(row=18, column=2).value,
        "deal_count": int(ws_pivot.cell(row=18, column=3).value or 0),
        "acv_usd": round(float(ws_pivot.cell(row=18, column=4).value or 0.0), 2)
    }

    # 5. Evaluate Subtotal (row 19)
    subtotal_count = sum(item["deal_count"] for item in ut30_values.values())
    subtotal_acv = round(sum(item["total_acv"] for item in ut30_values.values()), 2)
    subtotal_avg = round(subtotal_acv / subtotal_count, 2) if subtotal_count > 0 else 0.0
    subtotal = {
        "deal_count": subtotal_count,
        "total_acv": subtotal_acv,
        "avg_acv": subtotal_avg
    }

    # 6. Evaluate Grand Total (row 20)
    grand_count = subtotal_count + manual_adj["deal_count"]
    grand_acv = round(subtotal_acv + manual_adj["acv_usd"], 2)
    grand_avg = round(grand_acv / grand_count, 2) if grand_count > 0 else 0.0
    grand_total = {
        "deal_count": grand_count,
        "total_acv": grand_acv,
        "avg_acv": grand_avg
    }

    return {
        "ut30_values": ut30_values,
        "manual_adj": manual_adj,
        "subtotal": subtotal,
        "grand_total": grand_total
    }

def test_forecast_engine_parity_with_legacy_workbook(tmp_path):
    """
    Verify that the Forecast Engine matches the evaluated values from the legacy Excel workbook:
    - Grand Total ACV and Deal Count
    - Subtotal ACV and Deal Count
    - Every UT30 product line ACV to the exact cent.
    """
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    legacy_path = os.path.join(base_dir, "legacy", "Daily_Forecast_LEGACY.xlsx")
    snapshot_path = todays_snapshot_path(base_dir)

    if not os.path.exists(legacy_path):
        build_legacy_workbook(csv_path=snapshot_path, output_path=legacy_path)

    legacy_data = load_legacy_workbook_pivot_values(legacy_path)

    df_raw = pd.read_csv(snapshot_path)
    result_df, summary = compute_forecast_from_df(df_raw)

    # 1. Grand Total Parity (Includes manual adjustment)
    assert summary["grand_total_acv"] == legacy_data["grand_total"]["total_acv"], (
        f"Grand Total ACV mismatch: Engine=${summary['grand_total_acv']} vs Legacy=${legacy_data['grand_total']['total_acv']}"
    )
    assert summary["deal_count"] == legacy_data["grand_total"]["deal_count"], (
        f"Grand Total Deal Count mismatch: Engine={summary['deal_count']} vs Legacy={legacy_data['grand_total']['deal_count']}"
    )
    assert summary["avg_deal_size"] == legacy_data["grand_total"]["avg_acv"]

    # 2. Subtotal Parity (Extract only)
    assert summary["subtotal_acv"] == legacy_data["subtotal"]["total_acv"]
    assert summary["subtotal_deal_count"] == legacy_data["subtotal"]["deal_count"]

    # 3. UT30 Product Line Totals (to the cent)
    eng_breakdown = {item["ut_level_30_dscr"]: item for item in summary["ut30_breakdown"]}
    for prod_name, leg_item in legacy_data["ut30_values"].items():
        assert prod_name in eng_breakdown, f"Product {prod_name} missing from engine output"
        eng_item = eng_breakdown[prod_name]
        assert eng_item["total_acv"] == leg_item["total_acv"], (
            f"UT30 {prod_name} ACV mismatch: Engine=${eng_item['total_acv']} vs Legacy=${leg_item['total_acv']}"
        )
        assert eng_item["deal_count"] == leg_item["deal_count"], (
            f"UT30 {prod_name} deal count mismatch: Engine={eng_item['deal_count']} vs Legacy={leg_item['deal_count']}"
        )

def test_parity_fails_if_excluded_customers_rule_removed(monkeypatch):
    """
    Verify test fails if the 'excluded_customers' rule is removed.
    """
    import app.config
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    snapshot_path = todays_snapshot_path(base_dir)
    legacy_path = os.path.join(base_dir, "legacy", "Daily_Forecast_LEGACY.xlsx")
    legacy_data = load_legacy_workbook_pivot_values(legacy_path)

    # Inject a test record for the excluded customer into the raw data
    df_raw = pd.read_csv(snapshot_path)
    test_record = {
        "ut_level_30_dscr": "watsonx.data Core",
        "ut_level_15_dscr": "Data Platform",
        "coverage_id": "COV-00001",
        "signing_date": get_as_of_date().isoformat(),
        "sales_order_num": "SO-EXCLUDED-01",
        "customer_name": "Acme Test Enterprise 999",
        "saas_acv_usd_at_plan": 50000.0,
        "country": "United States",
        "client_sub_type": "Select T Activate",
        "signing_category": "New Signings",
    }
    df_with_excluded = pd.concat([df_raw, pd.DataFrame([test_record])], ignore_index=True)

    # Mock rules WITHOUT excluded_customers
    original_rules = app.config.get_rules_config()
    broken_rules = dict(original_rules)
    broken_rules["excluded_customers"] = []
    monkeypatch.setattr(app.config, "get_rules_config", lambda *args, **kwargs: broken_rules)

    _, summary = compute_forecast_from_df(df_with_excluded)
    # The summary will include the $50k deal, causing parity to fail
    assert summary["subtotal_acv"] != legacy_data["subtotal"]["total_acv"], "Engine should fail parity when excluded_customers is absent"

def test_parity_fails_if_manual_adjustments_rule_removed(monkeypatch):
    """
    Verify test fails if the 'manual_adjustments' rule is removed.
    """
    import app.config
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    snapshot_path = todays_snapshot_path(base_dir)
    legacy_path = os.path.join(base_dir, "legacy", "Daily_Forecast_LEGACY.xlsx")
    legacy_data = load_legacy_workbook_pivot_values(legacy_path)

    df_raw = pd.read_csv(snapshot_path)

    # Mock rules WITHOUT manual_adjustments
    original_rules = app.config.get_rules_config()
    broken_rules = dict(original_rules)
    broken_rules["manual_adjustments"] = []
    monkeypatch.setattr(app.config, "get_rules_config", lambda *args, **kwargs: broken_rules)

    _, summary = compute_forecast_from_df(df_raw)
    # Grand total will be missing the $150,000.00 adjustment
    assert summary["grand_total_acv"] != legacy_data["grand_total"]["total_acv"], "Engine should fail parity when manual_adjustments is absent"
    assert summary["grand_total_acv"] == legacy_data["subtotal"]["total_acv"]
